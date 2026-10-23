# -*- coding: utf-8 -*-
"""校验生成的立项书 docx：结构、段落数、表格、字体、页边距。"""
import sys
from docx import Document
from docx.shared import Pt

path = r"L:\WebLearn\ai-go-desktop\docs\时光围棋ShiGuang-Go_创业立项书.docx"
doc = Document(path)

paras = doc.paragraphs
heads = [(p.style.name, p.text) for p in paras if p.style.name.startswith("Heading")]
print("段落总数:", len(paras))
print("标题数:", len(heads))
print("表格数:", len(doc.tables))
for s, t in heads:
    print(f"  {s:10s} | {t}")

print("\n--- 表格规模 ---")
for i, t in enumerate(doc.tables):
    print(f"  表{i+1}: {len(t.rows)} 行 x {len(t.columns)} 列 | 首行: {t.rows[0].cells[0].text}")

sec = doc.sections[0]
print("\n--- 页面 ---")
print(f"  纸张: {sec.page_width.cm:.1f} x {sec.page_height.cm:.1f} cm")
print(f"  页边距 上{sec.top_margin.cm:.1f} 下{sec.bottom_margin.cm:.1f} "
      f"左{sec.left_margin.cm:.1f} 右{sec.right_margin.cm:.1f}")

# 中文字体检查
n = doc.styles["Normal"]
from docx.oxml.ns import qn
ea = n._element.rPr.rFonts.get(qn("w:eastAsia"))
print(f"\n正文样式: 中文={ea} 西文={n.font.name} 字号={n.font.size.pt if n.font.size else 'inherit'}")

# 空段落/异常检查
empty = sum(1 for p in paras if not p.text.strip())
print(f"\n空段落数: {empty}（封面留白属正常）")

# 关键内容抽查
full = "\n".join(p.text for p in paras)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            full += "\n" + c.text
checks = ["待核实", "1.06 GB", "15,995", "4,925", "29 档", "sm89", "买断", "P0 主攻"]
print("\n--- 关键内容抽查 ---")
for c in checks:
    print(f"  {'OK ' if c in full else 'MISS'} {c}")

print(f"\n总字符数(含表格): {len(full)}")
